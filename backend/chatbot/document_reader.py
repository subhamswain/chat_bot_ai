import csv
import os

import pymupdf
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation


# ============================================================
# PDF
# ============================================================

def read_pdf(file_path):
    """
    Extract text from PDF while preserving page numbers.
    """

    pdf = pymupdf.open(file_path)

    pages = []

    try:

        for page_number, page in enumerate(pdf, start=1):

            text = page.get_text("text")

            pages.append(
                f"--- Page {page_number} ---\n"
                f"{text.strip()}"
            )

    finally:

        pdf.close()

    return "\n\n".join(pages)


# ============================================================
# DOCX
# ============================================================

def read_docx(file_path):
    """
    Extract paragraphs and tables from Word documents.
    """

    document = DocxDocument(file_path)

    content = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            content.append(text)

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table_index, table in enumerate(
        document.tables,
        start=1
    ):

        content.append(
            f"\n--- Table {table_index} ---"
        )

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            content.append(
                " | ".join(cells)
            )

    return "\n".join(content)


# ============================================================
# EXCEL
# ============================================================

def read_excel(file_path):
    """
    Extract all sheets, rows and columns from Excel.
    """

    workbook = load_workbook(
        filename=file_path,
        data_only=True
    )

    content = []

    for sheet in workbook.worksheets:

        content.append(
            f"\n--- Sheet: {sheet.title} ---"
        )

        for row in sheet.iter_rows(
            values_only=True
        ):

            values = []

            for value in row:

                if value is None:
                    values.append("")
                else:
                    values.append(str(value))

            # Skip completely empty rows
            if any(
                value.strip()
                for value in values
            ):

                content.append(
                    " | ".join(values)
                )

    return "\n".join(content)


# ============================================================
# POWERPOINT
# ============================================================

def read_pptx(file_path):
    """
    Extract text from all PowerPoint slides.
    """

    presentation = Presentation(
        file_path
    )

    content = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        content.append(
            f"\n--- Slide {slide_number} ---"
        )

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:

                    content.append(text)

    return "\n".join(content)


# ============================================================
# TXT
# ============================================================

def read_txt(file_path):
    """
    Read normal text files.
    """

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1",
    ]

    for encoding in encodings:

        try:

            with open(
                file_path,
                "r",
                encoding=encoding
            ) as file:

                return file.read()

        except UnicodeDecodeError:

            continue

    raise ValueError(
        "Could not decode the text file."
    )


# ============================================================
# CSV
# ============================================================

def read_csv(file_path):
    """
    Extract CSV rows and columns.
    """

    content = []

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1",
    ]

    for encoding in encodings:

        try:

            with open(
                file_path,
                "r",
                encoding=encoding,
                newline=""
            ) as file:

                reader = csv.reader(file)

                for row in reader:

                    if any(
                        str(value).strip()
                        for value in row
                    ):

                        content.append(
                            " | ".join(
                                str(value)
                                for value in row
                            )
                        )

            return "\n".join(content)

        except UnicodeDecodeError:

            content = []
            continue

    raise ValueError(
        "Could not decode the CSV file."
    )


# ============================================================
# UNIVERSAL DOCUMENT READER
# ============================================================

def read_document(file_path):
    """
    Automatically detect file type and extract content.
    """

    extension = os.path.splitext(
        file_path
    )[1].lower()

    print(
        "DOCUMENT EXTENSION:",
        extension
    )

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == ".pdf":

        return read_pdf(
            file_path
        )

    # --------------------------------------------------------
    # Word
    # --------------------------------------------------------

    if extension == ".docx":

        return read_docx(
            file_path
        )

    # --------------------------------------------------------
    # Excel
    # --------------------------------------------------------

    if extension in [
        ".xlsx",
        ".xlsm",
    ]:

        return read_excel(
            file_path
        )

    # --------------------------------------------------------
    # PowerPoint
    # --------------------------------------------------------

    if extension == ".pptx":

        return read_pptx(
            file_path
        )

    # --------------------------------------------------------
    # Text
    # --------------------------------------------------------

    if extension == ".txt":

        return read_txt(
            file_path
        )

    # --------------------------------------------------------
    # CSV
    # --------------------------------------------------------

    if extension == ".csv":

        return read_csv(
            file_path
        )

    # --------------------------------------------------------
    # Unsupported
    # --------------------------------------------------------

    raise ValueError(
        f"Unsupported file type: {extension}"
    )